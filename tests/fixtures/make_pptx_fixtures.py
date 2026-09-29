"""Generates .pptx fixtures for tests/pptx-ops.test.mts (requires python-pptx).

usage: python3 make_pptx_fixtures.py OUT_DIR
"""
import copy
import struct
import sys
import zlib
from io import BytesIO
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
SECTION_URI = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"


def tiny_png(rgb):
    def chunk(kind, data):
        c = struct.pack(">I", len(data)) + kind + data
        return c + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)

    raw = b"".join(b"\x00" + bytes(rgb) * 8 for _ in range(8))
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 8, 8, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def add_sections(prs, groups):
    """groups: list of (name, [slide indexes])"""
    pres = prs.part._element
    ids = [s.get("id") for s in pres.find(qn("p:sldIdLst"))]
    ext_lst = pres.find(qn("p:extLst"))
    if ext_lst is None:
        ext_lst = etree.SubElement(pres, qn("p:extLst"))
    ext = etree.SubElement(ext_lst, qn("p:ext"))
    ext.set("uri", SECTION_URI)
    sl = etree.SubElement(ext, "{%s}sectionLst" % P14, nsmap={"p14": P14})
    for k, (name, idxs) in enumerate(groups):
        sec = etree.SubElement(sl, "{%s}section" % P14)
        sec.set("name", name)
        sec.set("id", "{0000000%d-0000-0000-0000-000000000000}" % (k + 1))
        lst = etree.SubElement(sec, "{%s}sldIdLst" % P14)
        for i in idxs:
            e = etree.SubElement(lst, "{%s}sldId" % P14)
            e.set("id", ids[i])


def add_custom_show(prs, name, slide_indexes):
    pres = prs.part._element
    rids = [s.get(qn("r:id")) for s in pres.find(qn("p:sldIdLst"))]
    lst = etree.Element(qn("p:custShowLst"))
    show = etree.SubElement(lst, qn("p:custShow"))
    show.set("name", name)
    show.set("id", "0")
    sld_lst = etree.SubElement(show, qn("p:sldLst"))
    for i in slide_indexes:
        e = etree.SubElement(sld_lst, qn("p:sld"))
        e.set(qn("r:id"), rids[i])
    # custShowLst comes after notesSz (and before photoAlbum/custDataLst/kinsoku/defaultTextStyle...)
    notes_sz = pres.find(qn("p:notesSz"))
    notes_sz.addnext(lst)


def rich_deck(path):
    prs = Presentation()
    title_layout = prs.slide_layouts[1]  # Title and Content
    blank = prs.slide_layouts[6]
    png = BytesIO(tiny_png((200, 40, 90)))

    # Slide 1: title + multi-run formatted paragraphs + shared picture + link to slide 4
    s1 = prs.slides.add_slide(title_layout)
    s1.shapes.title.text = "Alpha"
    body = s1.placeholders[1].text_frame
    body.text = ""
    p0 = body.paragraphs[0]
    r = p0.add_run()
    r.text = "Bold red "
    r.font.bold = True
    r.font.size = Pt(28)
    r.font.color.rgb = RGBColor(0xC0, 0x10, 0x20)
    r2 = p0.add_run()
    r2.text = "plain tail"
    p1 = body.add_paragraph()
    p1.level = 1
    r3 = p1.add_run()
    r3.text = "Italic second"
    r3.font.italic = True
    r3.font.size = Pt(18)
    p2 = body.add_paragraph()
    p2.text = "Line one\vLine two"  # a:br between two runs
    for run in p2.runs:
        run.font.underline = True
    pic = s1.shapes.add_picture(png, Inches(8), Inches(0.2), Inches(1), Inches(1))

    # Slide 2: notes + group with text shapes + a field
    s2 = prs.slides.add_slide(blank)
    grp = s2.shapes.add_group_shape()
    g1 = grp.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    g1.text_frame.text = "Group child one"
    g2 = grp.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1))
    g2.text_frame.text = "Group child two"
    tb = s2.shapes.add_textbox(Inches(5), Inches(1), Inches(3), Inches(1))
    tb.text_frame.text = "Slide "
    p = tb.text_frame.paragraphs[0]
    fld = etree.SubElement(p._p, qn("a:fld"))
    fld.set("id", "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}")
    fld.set("type", "slidenum")
    frpr = etree.SubElement(fld, qn("a:rPr"))
    frpr.set("lang", "en-US")
    ft = etree.SubElement(fld, qn("a:t"))
    ft.text = "2"
    s2.notes_slide.notes_text_frame.text = "Notes for Beta"

    # Slide 3: table + same picture (shared media) + notes
    s3 = prs.slides.add_slide(title_layout)
    s3.shapes.title.text = "Gamma"
    tbl = s3.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1.5)).table
    tbl.cell(0, 0).text = "H1"
    tbl.cell(0, 1).text = "H2"
    tbl.cell(1, 0).text = "c10"
    tbl.cell(1, 1).text = "c11 & <x>"
    png.seek(0)
    s3.shapes.add_picture(png, Inches(8), Inches(0.2), Inches(1), Inches(1))
    s3.notes_slide.notes_text_frame.text = "Notes for Gamma"

    # Slide 4..5
    s4 = prs.slides.add_slide(title_layout)
    s4.shapes.title.text = "Delta"
    chart_data = CategoryChartData()
    chart_data.categories = ["A", "B"]
    chart_data.add_series("S", (1, 2))
    s4.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(4), Inches(3), chart_data)
    s5 = prs.slides.add_slide(title_layout)
    s5.shapes.title.text = "Epsilon"
    # A text box wrapped in mc:AlternateContent (Choice + Fallback share the id)
    alt_box = s5.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1))
    alt_box.text_frame.text = "Alt text"
    sp = alt_box._element
    MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    alt = etree.Element("{%s}AlternateContent" % MC, nsmap={"mc": MC, "p14": P14})
    sp.addprevious(alt)
    choice = etree.SubElement(alt, "{%s}Choice" % MC)
    choice.set("Requires", "p14")
    fallback = etree.SubElement(alt, "{%s}Fallback" % MC)
    choice.append(copy.deepcopy(sp))
    fallback.append(sp)

    # Hyperlink on slide 1's picture to slide 4
    pic.click_action.target_slide = s4

    add_sections(prs, [("Intro", [0, 1]), ("Middle", [2, 3]), ("Empty", []), ("End", [4])])
    add_custom_show(prs, "Short", [0, 2, 3])
    prs.save(path)


def plain_deck(path, n=3):
    prs = Presentation()
    for i in range(n):
        s = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
        s.shapes.title.text = "Plain %d" % (i + 1)
    prs.save(path)


if __name__ == "__main__":
    out = Path(sys.argv[1])
    out.mkdir(parents=True, exist_ok=True)
    rich_deck(str(out / "rich.pptx"))
    plain_deck(str(out / "plain.pptx"))
    print("ok")
