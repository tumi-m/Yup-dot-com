"""Dumps what python-pptx sees in each .pptx as JSON (used by tests/pptx-ops.test.mts).

usage: python3 inspect_pptx.py FILE...
"""
import json
import sys

from pptx import Presentation
from pptx.oxml.ns import qn

P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"


def shape_texts(shapes, out):
    for sh in shapes:
        if sh.shape_type is not None and sh.shape_type == 6:  # GROUP
            shape_texts(sh.shapes, out)
            continue
        if getattr(sh, "has_table", False) and sh.has_table:
            for r, row in enumerate(sh.table.rows):
                for c, cell in enumerate(row.cells):
                    out.append({"id": "%d:%d:%d" % (sh.shape_id, r, c), "text": cell.text})
            continue
        if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
            out.append({"id": str(sh.shape_id), "text": sh.text_frame.text})


def inspect(path):
    try:
        prs = Presentation(path)
    except Exception as e:  # noqa: BLE001
        return {"error": repr(e)}
    pres = prs.part._element
    slides = []
    for s in prs.slides:
        texts = []
        shape_texts(s.shapes, texts)
        title = s.shapes.title.text if s.shapes.title is not None else None
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else None
        slides.append({"part": str(s.part.partname), "title": title, "texts": texts, "notes": notes})
    ids = [el.get("id") for el in pres.find(qn("p:sldIdLst"))]
    rids = [el.get(qn("r:id")) for el in pres.find(qn("p:sldIdLst"))]
    sections = []
    for sec in pres.iter("{%s}section" % P14):
        lst = sec.find("{%s}sldIdLst" % P14)
        sections.append(
            {"name": sec.get("name"), "ids": [e.get("id") for e in (lst if lst is not None else [])]}
        )
    shows = []
    for show in pres.iter(qn("p:custShow")):
        shows.append([e.get(qn("r:id")) for e in show.iter(qn("p:sld"))])
    return {"slides": slides, "sldIds": ids, "rIds": rids, "sections": sections, "shows": shows}


if __name__ == "__main__":
    print(json.dumps({p: inspect(p) for p in sys.argv[1:]}))
